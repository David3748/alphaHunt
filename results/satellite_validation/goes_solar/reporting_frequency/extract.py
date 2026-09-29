#!/usr/bin/env python3
"""Extract two plants' EIA-923 reporting metadata and annual generation labels.

Annual label ingestion was explicitly requested after the metadata-only audit.
This code fits no forecast and selects no plant based on generation outcomes.
"""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
CACHE = ROOT / "work/eia_frequency_audit"


def inspect(year):
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / f"f923_{year}.zip"
    url = f"https://www.eia.gov/electricity/data/eia923/{'xls' if year == 2025 else 'archive/xls'}/f923_{year}.zip"
    metadata = {"year": year, "url": url}
    if not path.exists():
        response = requests.get(url, timeout=180)
        response.raise_for_status()
        path.write_bytes(response.content)
        metadata["http_last_modified"] = response.headers.get("Last-Modified")
        metadata["retrieved_at_utc"] = datetime.now(timezone.utc).isoformat()
        path.with_suffix(".json").write_text(json.dumps(metadata, indent=2) + "\n")
    else:
        metadata.update(json.loads(path.with_suffix(".json").read_text()))
    metadata["zip_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    metadata["zip_bytes"] = path.stat().st_size
    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith(".xlsx") and "2_3_4_5" in n]
        if len(names) != 1:
            raise ValueError(f"Ambiguous generation workbook: {names}")
        name = names[0]
        metadata["workbook"] = name
        raw = archive.read(name)
        metadata["workbook_sha256"] = hashlib.sha256(raw).hexdigest()
        book = openpyxl.load_workbook(archive.open(name), read_only=True, data_only=True)
        sheets = [s for s in book.sheetnames if "Plant Frame" in s and "Puerto Rico" not in s]
        if len(sheets) != 1:
            raise ValueError(f"Ambiguous plant metadata sheet: {sheets}")
        metadata["sheet"] = sheets[0]
        header = None
        rows = []
        for index, row in enumerate(book[sheets[0]].iter_rows(values_only=True), start=1):
            normalized = [" ".join(str(x).split()).lower() if x is not None else "" for x in row]
            if "plant id" in normalized and ("respondent frequency" in normalized or "reporting frequency" in normalized):
                header = normalized
                continue
            if header is None:
                continue
            values = dict(zip(header, row))
            if values["plant id"] in {57439, 57695}:
                rows.append({"year": year, "plant_id": values["plant id"],
                             "plant_name": values["plant name"],
                             "respondent_frequency": values.get("respondent frequency", values.get("reporting frequency")),
                             "sheet_row": index, "url": url, "zip_sha256": metadata["zip_sha256"]})
        if len(rows) != 2:
            raise ValueError(f"Expected exactly two plant metadata rows in {year}: {rows}")
        annual = []
        header = None
        for index, row in enumerate(book["Page 1 Generation and Fuel Data"].iter_rows(values_only=True), start=1):
            normalized = [" ".join(str(x).split()).lower() if x is not None else "" for x in row]
            if "plant id" in normalized and "reported prime mover" in normalized:
                header = normalized
                continue
            if header is None:
                continue
            values = dict(zip(header, row))
            if values["plant id"] not in {57439, 57695}:
                continue
            if values["reported prime mover"] != "PV" or values["reported fuel type code"] != "SUN":
                continue
            record = next(r for r in rows if r["plant_id"] == values["plant id"])
            total = float(values["net generation (megawatthours)"])
            months = [float(v) for k, v in values.items() if k.startswith("netgen ")]
            if len(months) != 12:
                raise ValueError("Expected 12 monthly allocation/reported fields")
            if abs(total - sum(months)) > 1:
                raise ValueError("Annual net generation and monthly sum differ by over 1 MWh")
            annual.append({**record, "annual_generation_mwh": total,
                           "monthly_sum_mwh": sum(months), "generation_sheet_row": index,
                           "generation_source_field": "Net Generation (Megawatthours)",
                           "workbook_sha256": metadata["workbook_sha256"]})
        if len(annual) != 2:
            raise ValueError(f"Expected exactly two annual PV/SUN generation rows in {year}")
        layouts = [s for s in book.sheetnames if "File Layout" in s]
        definitions = []
        for sheet in layouts:
            iterator = list(book[sheet].iter_rows(values_only=True))
            for i, row in enumerate(iterator):
                if any(isinstance(x, str) and x.lower() == "respondent frequency" for x in row):
                    definitions = [[str(x) if x is not None else None for x in r] for r in iterator[i:i+4]]
                    break
        metadata["frequency_definition_rows"] = definitions
        book.close()
    print(year, [(r["plant_name"], r["respondent_frequency"]) for r in rows], flush=True)
    return rows, metadata, annual


if __name__ == "__main__":
    with ThreadPoolExecutor(max_workers=3) as pool:
        results = list(pool.map(inspect, range(2015, 2026)))
    pd.DataFrame([row for rows, _, _ in results for row in rows]).to_csv(HERE / "plant_reporting_frequency.csv", index=False)
    pd.DataFrame([row for _, _, annual in results for row in annual]).to_csv(HERE / "annual_generation.csv", index=False)
    manifest = {"created_at_utc": datetime.now(timezone.utc).isoformat(),
                "method": "Read Page 6 Plant Frame, select plant IDs 57439 and 57695; read frequency definitions from File Layout. Subsequently ingest annual Net Generation (Megawatthours) for exact PV/SUN rows in Page 1, verify monthly sums within 1 MWh. No model fitting. Current final archives, not original monthly snapshots.",
                "sources": [metadata for _, metadata, _ in results]}
    (HERE / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
